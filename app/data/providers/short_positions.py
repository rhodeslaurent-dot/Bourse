"""docs/23 S2 — net short positions published by market authorities (EU regulation 236/2012: public
from 0.5 % of the share capital). Each authority publishes its own CSV / Excel file; columns are
recognised by their header words (French, English, German, Dutch, Spanish, Swedish…), never by
position, and a file whose columns cannot all be recognised is refused rather than guessed."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timedelta

import httpx

from app.data.base import ProviderError

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36"
_ISIN = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$")

# header words per field, accent-free and lower-case (matched as substrings)
HEADERS: dict[str, tuple[str, ...]] = {
    "isin": ("isin",),
    "holder": (
        "detenteur", "holder", "positionsinhaber", "inhaber", "innehavare", "tenedor", "titular", "houder",
        "meldingsplichtige", "name of the position", "position holder",
    ),
    "issuer": ("emetteur", "issuer", "emittent", "emisor", "uitgevende", "naam van de uitgevende", "name of the share"),
    "pct": (
        "position courte", "net short position", "netto-leerverkaufsposition", "leerverkaufsposition", "posicion",
        "positie", "short position", "positionens storlek", "position in %", "position (%)", "netto short",
        "net short", "%", "position",
    ),
    "date": ("date", "datum", "fecha", "dato"),
}  # fmt: skip


@dataclass(frozen=True)
class ShortRow:
    isin: str
    issuer: str | None
    holder: str
    pct: float  # fraction (0.0055 = 0.55 %)
    position_date: date


def _norm(s: object) -> str:
    text = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return " ".join(text.split())


def map_columns(header: list[object]) -> dict[str, int] | None:
    """Field → column index from a header row; None when isin, holder, pct or date is missing. The ISIN
    column is matched first, then the most specific words, so « date de la position » never takes the
    position column."""
    cells = [_norm(h) for h in header]
    out: dict[str, int] = {}
    for field in ("isin", "holder", "issuer", "date", "pct"):
        for i, c in enumerate(cells):
            if i in out.values() or not c:
                continue
            words = HEADERS[field]
            if field == "pct" and ("date" in c or "datum" in c or "fecha" in c):
                continue
            if any(w in c for w in words):
                out[field] = i
                break
    return out if all(k in out for k in ("isin", "holder", "pct", "date")) else None


def parse_pct(value: object) -> float | None:
    """Published in percent of the capital (0,55 → 0.55 %): returned as a fraction. A value already
    below 0.1 is taken as a fraction (no public position is under 0.5 %, except the « below » notice)."""
    if isinstance(value, int | float):
        v = float(value)
    else:
        txt = str(value or "").strip().replace("%", "").replace(" ", "").replace(" ", "").replace(",", ".")
        try:
            v = float(txt)
        except ValueError:
            return None
    return v / 100 if v >= 0.1 else v


def parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, int | float) and 20000 < value < 80000:  # Excel serial day
        return date(1899, 12, 30) + timedelta(days=int(value))
    txt = str(value or "").strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(txt, fmt).date()
        except ValueError:
            continue
    return None


def rows_from_table(table: list[list[object]]) -> list[ShortRow]:
    """Find the header row (the first one whose columns are all recognised, within the first 40 rows)
    and read the positions below it; rows without a valid ISIN, percentage or date are skipped."""
    for k, header in enumerate(table[:40]):
        cols = map_columns(header)
        if cols is None:
            continue
        out = []
        for raw in table[k + 1 :]:

            def cell(f: str, raw: list[object] = raw, cols: dict[str, int] = cols) -> object:
                return raw[cols[f]] if f in cols and cols[f] < len(raw) else None

            isin = str(cell("isin") or "").strip().upper()
            pct, day = parse_pct(cell("pct")), parse_date(cell("date"))
            if not _ISIN.match(isin) or pct is None or day is None:
                continue
            issuer = str(cell("issuer")).strip() if cell("issuer") else None
            out.append(ShortRow(isin, issuer, str(cell("holder") or "").strip()[:200] or "?", pct, day))
        return out
    raise ProviderError("colonnes non reconnues (ISIN, détenteur, position, date) : fichier refusé")


def read_table(content: bytes, name: str = "") -> list[list[object]]:
    """CSV (delimiter sniffed, UTF-8 or Latin-1) or Excel (.xlsx) into rows of cells."""
    if content[:2] == b"PK" or name.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise ProviderError("openpyxl non installé : fichier Excel illisible") from exc
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows: list[list[object]] = []
        for ws in wb.worksheets:
            rows += [list(r) for r in ws.iter_rows(values_only=True)]
        return rows
    for enc in ("utf-8-sig", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4000]
    delim = max((";", ",", "\t", "|"), key=sample.count)
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]


def resolve_datagouv(query: str, client: httpx.Client) -> str:
    """URL of the latest CSV/Excel resource of the data.gouv.fr dataset matching ``query`` (AMF open
    data, daily update): the address is read from the official catalogue API, never hard-coded."""
    r = client.get("https://www.data.gouv.fr/api/1/datasets/", params={"q": query, "page_size": 10})
    if r.status_code >= 400:
        raise ProviderError(f"data.gouv.fr {r.status_code}")
    words = [_norm(w) for w in query.split()]
    for ds in r.json().get("data", []):
        if not all(w in _norm(ds.get("title")) for w in words):
            continue
        res = [x for x in ds.get("resources", []) if _norm(x.get("format")) in ("csv", "xlsx", "xls")]
        res.sort(key=lambda x: str(x.get("last_modified") or ""), reverse=True)
        if res:
            return str(res[0]["url"])
    raise ProviderError(f"aucun jeu de données data.gouv.fr pour « {query} »")


def fetch(source: dict, client: httpx.Client | None = None) -> tuple[list[ShortRow], str]:
    """(positions, URL read) for one configured authority; ``url: "datagouv:<words>"`` is resolved first."""
    c = client or httpx.Client(timeout=60, follow_redirects=True, headers={"User-Agent": UA})
    url = str(source.get("url") or "")
    if not url or "renseigner" in url:
        raise ProviderError("adresse à renseigner")
    if url.startswith("datagouv:"):
        url = resolve_datagouv(url.split(":", 1)[1], c)
    r = c.get(url)
    if r.status_code >= 400:
        raise ProviderError(f"HTTP {r.status_code}")
    return rows_from_table(read_table(r.content, url)), url
